"""Паспорт коллекции (MEM-23, MEM-77).

Берёт ссылки из links_*.txt (в любом виде — скрипт сам вытащит URL), убирает дубли,
подтягивает метаданные через yt-dlp и пишет passport.xlsx + passport.csv.
Ручные столбцы (речь, язык, текст в кадре, тема, звук, комментарий) при повторном
запуске сохраняются. Лист «Сводка» считает критерии приёмки формулами.

Запуск:  python passport.py            (все ссылки)
         python passport.py --offline  (без yt-dlp, только список и дубли)
"""

import csv
import glob
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "passport.xlsx")
CSV = os.path.join(HERE, "passport.csv")
OWNERS = {"anastasia": "Анастасия", "anastasia_vk": "Анастасия", "nika": "Ника"}
TOPICS = [
    "Юмор и мемы",
    "Еда и рецепты",
    "Красота и уход",
    "Мода и покупки",
    "Учёба и IT",
    "Спорт и здоровье",
    "Путешествия и места",
    "Психология и отношения",
    "Музыка и танцы",
    "Лайфхаки и быт",
    "Кино и сериалы",
    "Животные",
    "Общество и политика",
    "Другое",
]

AUTO = [
    "№",
    "ID",
    "Платформа",
    "Ссылка",
    "Чьё «Избранное»",
    "Автор",
    "Название / подпись",
    "Длительность, с",
    "Описание есть",
    "Субтитры платформы",
    "Статус ссылки",
]
MANUAL = ["Речь", "Язык", "Текст в кадре", "Тема", "Качество звука", "Комментарий"]
COLS = AUTO + MANUAL

URL_RE = re.compile(r'https?://[^\s"\'<>]+')
# Ролики собственных аккаунтов команды в коллекцию не берём
EXCLUDE_AUTHORS = {"@evv.zt"}


def canon(url):
    """(platform, id) для дедупликации или None, если ссылка не TikTok/VK."""
    u = url.split("?")[0].rstrip("/")
    m = re.search(r"tiktok\.com/@[^/]+/(?:video|photo)/(\d+)", u)
    if m:
        return "TikTok", m.group(1)
    if re.search(r"(vm|vt)\.tiktok\.com/\w+", u):
        return "TikTok", "short:" + u.rsplit("/", 1)[-1]
    m = re.search(r"(?:vk\.com|vkvideo\.ru|vk\.ru)/(?:.*?)(clip|video)(-?\d+_\d+)", u)
    if m:
        return "VK", m.group(2)
    return None


def read_links():
    rows = []
    for path in sorted(glob.glob(os.path.join(HERE, "links_*.txt"))):
        who = OWNERS.get(os.path.basename(path)[6:-4], os.path.basename(path)[6:-4])
        with open(path, encoding="utf-8") as f:
            for url in URL_RE.findall(f.read()):
                rows.append((url, who))
    return rows


def ytdlp_info(url):
    try:
        import yt_dlp
    except ImportError:
        return None, "yt-dlp не установлен"
    opts = {"quiet": True, "no_warnings": True, "skip_download": True, "socket_timeout": 20}
    try:
        with yt_dlp.YoutubeDL(opts) as y:
            return y.extract_info(url, download=False), "доступна"
    except Exception as e:
        msg = str(e).split("\n")[0]
        return None, "ошибка: " + msg[:120]


def tiktok_oembed(url):
    """Запасной путь, если yt-dlp не достучался до TikTok (из России): автор, подпись, доступность."""
    import json
    import urllib.parse
    import urllib.request

    try:
        q = "https://www.tiktok.com/oembed?url=" + urllib.parse.quote(url, safe="")
        with urllib.request.urlopen(
            urllib.request.Request(q, headers={"User-Agent": "Mozilla/5.0"}), timeout=20
        ) as r:
            j = json.loads(r.read().decode("utf-8"))
        author = (j.get("author_url") or "").rsplit("@", 1)[-1] or j.get("author_name", "")
        return {
            "uploader": "@" + author if author else "",
            "title": j.get("title") or "",
        }, "доступна (oEmbed)"
    except Exception as e:
        return None, "ошибка: " + str(e)[:120]


def load_existing():
    if not os.path.exists(XLSX):
        return {}
    ws = load_workbook(XLSX)["Паспорт"]
    head = [c.value for c in ws[1]]
    keep = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(head, r, strict=False))
        if d.get("ID"):
            keep[(d["Платформа"], str(d["ID"]))] = d
    return keep


def main():
    offline = "--offline" in sys.argv
    old = load_existing()
    seen, rows, dups, bad = {}, [], [], []
    for url, who in read_links():
        c = canon(url)
        if not c:
            bad.append(url)
            continue
        if c in seen:
            dups.append(url)
            continue
        seen[c] = True
        d = dict(old.get(c, {}))
        d.update(
            {"ID": c[1], "Платформа": c[0], "Ссылка": url.split("?")[0], "Чьё «Избранное»": who}
        )
        if not offline and not str(d.get("Статус ссылки") or "").startswith("доступна"):
            if c[0] == "TikTok" and "--ytdlp" not in sys.argv:
                info, status = tiktok_oembed(url)  # yt-dlp из России до TikTok не достучивается
            else:
                info, status = ytdlp_info(url)
                if info is None and c[0] == "TikTok":
                    info, status = tiktok_oembed(url)
            d["Статус ссылки"] = status
            if info:
                if c[1].startswith("short:") and info.get("id"):
                    d["ID"] = str(info["id"])
                d["Автор"] = info.get("uploader") or info.get("channel") or ""
                d["Название / подпись"] = (info.get("title") or info.get("description") or "")[:200]
                d["Длительность, с"] = round(info["duration"]) if info.get("duration") else None
                cap = (info.get("description") or info.get("title") or "").strip()
                d["Описание есть"] = "да" if cap else "нет"
                if "subtitles" in info or "automatic_captions" in info:
                    d["Субтитры платформы"] = (
                        "да" if (info.get("subtitles") or info.get("automatic_captions")) else "нет"
                    )
            print(
                f"{len(rows) + 1:>3} {c[0]:6} {status[:40]}  {d.get('Название / подпись', '')[:50]}"
            )
        if d.get("Автор") in EXCLUDE_AUTHORS:
            print("    исключён: собственный аккаунт", d["Автор"])
            continue
        rows.append(d)
    write(rows)
    print(f"\nроликов: {len(rows)}, дублей отброшено: {len(dups)}, не TikTok/VK: {len(bad)}")
    for u in bad:
        print("  пропущено:", u)


def write(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Паспорт"
    ws.append(COLS)
    for i, d in enumerate(rows, 1):
        d["№"] = i
        ws.append([d.get(c) for c in COLS])
    n = len(rows) + 1
    hdr = PatternFill("solid", fgColor="E7E6E6")
    man = PatternFill("solid", fgColor="FFF2CC")
    for j, c in enumerate(COLS, 1):
        cell = ws.cell(1, j)
        cell.font = Font(bold=True)
        cell.fill = man if c in MANUAL else hdr
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = {
        "№": 5,
        "ID": 22,
        "Платформа": 10,
        "Ссылка": 45,
        "Чьё «Избранное»": 12,
        "Автор": 18,
        "Название / подпись": 50,
        "Длительность, с": 11,
        "Описание есть": 10,
        "Субтитры платформы": 11,
        "Статус ссылки": 18,
        "Речь": 8,
        "Язык": 8,
        "Текст в кадре": 10,
        "Тема": 22,
        "Качество звука": 13,
        "Комментарий": 30,
    }
    for j, c in enumerate(COLS, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = widths.get(c, 12)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLS)).column_letter}{max(n, 2)}"
    col = {c: ws.cell(1, j).column_letter for j, c in enumerate(COLS, 1)}
    # справочник тем
    ref = wb.create_sheet("Темы")
    for t in TOPICS:
        ref.append([t])

    def dv(letter, formula, allow_other=False):
        v = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=not allow_other,
            errorStyle="warning" if allow_other else "stop",
        )
        ws.add_data_validation(v)
        v.add(f"{letter}2:{letter}500")

    dv(col["Речь"], '"да,нет"')
    dv(col["Текст в кадре"], '"да,нет"')
    dv(col["Язык"], '"рус,англ,другой,нет речи"', True)
    dv(col["Качество звука"], '"хорошее,среднее,плохое,только музыка"')
    dv(col["Тема"], f"Темы!$A$1:$A${len(TOPICS)}", True)
    red = PatternFill("solid", fgColor="F8CBAD")
    ws.conditional_formatting.add(
        f"{col['Статус ссылки']}2:{col['Статус ссылки']}500",
        FormulaRule(formula=[f'LEFT({col["Статус ссылки"]}2,6)="ошибка"'], fill=red),
    )
    # сводка
    s = wb.create_sheet("Сводка", 0)
    P, R, T, TH, ST = (
        f"Паспорт!{col[k]}2:{col[k]}500"
        for k in ["Платформа", "Речь", "Текст в кадре", "Тема", "Статус ссылки"]
    )
    s.append(["Показатель", "Значение", "Критерий", "Выполнено"])
    s.append(
        [
            "Роликов в коллекции",
            "=COUNTA(Паспорт!A2:A500)",
            "90–110",
            '=IF(AND(B2>=90,B2<=110),"да","нет")',
        ]
    )
    s.append(["Из TikTok", f'=COUNTIF({P},"TikTok")', "", ""])
    s.append(["Из VK", f'=COUNTIF({P},"VK")', "", ""])
    s.append(["Доля TikTok", "=IF(B2=0,0,B3/B2)", "≥ 2/3", '=IF(B5>=2/3,"да","нет")'])
    s.append(["С речью", f'=COUNTIF({R},"да")', "", ""])
    s.append(
        [
            "Доля с речью или текстом в кадре",
            f'=IF(B2=0,0,SUMPRODUCT(--((({R})="да")+(({T})="да")>0))/B2)',
            "≥ 80 %",
            '=IF(B7>=0.8,"да","нет")',
        ]
    )
    s.append(["Доля без речи", f'=IF(B2=0,0,COUNTIF({R},"нет")/B2)', "посчитать", ""])
    s.append(["С текстом в кадре", f'=COUNTIF({T},"да")', "≥ 10", '=IF(B9>=10,"да","нет")'])
    s.append(
        [
            "Тем использовано",
            f'=SUMPRODUCT(({TH}<>"")/COUNTIF({TH},{TH}&""))',
            "≥ 5",
            '=IF(B10>=5,"да","нет")',
        ]
    )
    s.append(["Недоступных ссылок", f'=COUNTIF({ST},"ошибка*")', "0", '=IF(B11=0,"да","нет")'])
    s.append(["Не размечено (пустая «Речь»)", f"=B2-COUNTA({R})", "0", '=IF(B12=0,"да","нет")'])
    s.append([])
    s.append(["Распределение по темам", "Роликов", "Доля"])
    for i, t in enumerate(TOPICS):
        r = 15 + i
        s.append([t, f"=COUNTIF({TH},A{r})", f"=IF($B$2=0,0,B{r}/$B$2)"])
    for r in (5, 7, 8):
        s.cell(r, 2).number_format = "0%"
    for r in range(15, 15 + len(TOPICS)):
        s.cell(r, 3).number_format = "0%"
    for c in s[1] + s[14]:
        c.font = Font(bold=True)
    s.column_dimensions["A"].width = 32
    s.column_dimensions["B"].width = 12
    s.column_dimensions["C"].width = 12
    s.column_dimensions["D"].width = 12
    wb.save(XLSX)
    with open(CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        for d in rows:
            w.writerow([d.get(c) for c in COLS])
    print("сохранено:", XLSX)


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    main()
