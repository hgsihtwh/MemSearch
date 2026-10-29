"""Таблица ручной оценки ASR (MEM-16): отдельный лист на каждый ролик, фразы с таймкодами построчно.
Источник — results/ (выгрузка Colab-ноутбука asr_check.ipynb)."""

import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(HERE, "results")
res = json.load(open(os.path.join(RES, "results.json"), encoding="utf-8"))
by = {}
for r in res:
    by.setdefault(r["n"], {})[r["model"]] = r
clips = sorted(by)

HDR = PatternFill("solid", fgColor="D9D9D9")
MAN = PatternFill("solid", fgColor="FFF2CC")
EMPTY = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
F = Font(name="Calibri", size=12)
FB = Font(name="Calibri", size=12, bold=True)


def segments(n, model):
    p = os.path.join(RES, "transcripts", f"{n:02d}_{model}.txt")
    if not os.path.exists(p):
        return []
    out = []
    for line in open(p, encoding="utf-8").read().splitlines()[1:]:
        m = re.match(r"\[(\S+) – (\S+)\] (.*)", line)
        if m:
            out.append((m.group(1)[:5], m.group(3)))
    return out


def height(text, width_chars):
    lines = max(1, -(-len(text) // max(1, int(width_chars * 1.05))))
    return max(18, 16 * lines + 4)


wb = Workbook()
# --- Оценка (сводная, оценки подтягиваются с листов роликов)
ov = wb.active
ov.title = "Оценка"
ov.append(
    [
        "№",
        "Платформа",
        "Звук",
        "Тема",
        "Ролик",
        "Лист с транскриптом",
        "Длит., с",
        "Слов large-v3",
        "Оценка large-v3",
        "Оценка small",
        "Ошибки / комментарий",
    ]
)
for c in ov[1]:
    c.font = FB
    c.fill = HDR
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = BOX
for col, w in zip("ABCDEFGHIJK", [5, 10, 14, 34, 10, 14, 9, 10, 11, 11, 50], strict=True):
    ov.column_dimensions[col].width = w

dv_list = []
for i, n in enumerate(clips, 2):
    L, S = by[n].get("large-v3", {}), by[n].get("small", {})
    name = f"№{n}"
    sh = wb.create_sheet(name)
    # шапка листа ролика
    sh["A1"] = f"№{n} · {L.get('platform')} · звук: {L.get('sound')} · {L.get('topic')}"
    sh["A1"].font = Font(size=14, bold=True)
    sh["A2"] = "Открыть ролик"
    sh["A2"].hyperlink = L.get("url")
    sh["A2"].font = Font(size=12, color="0563C1", underline="single")
    sh["C2"] = "← к общему листу «Оценка»"
    sh["C2"].hyperlink = "#'Оценка'!A1"
    sh["C2"].font = Font(size=12, color="0563C1", underline="single")
    for r_, label in [
        (4, "Оценка large-v3 (2/1/0)"),
        (5, "Оценка small (2/1/0)"),
        (6, "Ошибки / комментарий"),
    ]:
        sh.cell(r_, 1, label).font = FB
        c = sh.cell(r_, 2)
        c.fill = MAN
        c.border = BOX
        c.font = F
        c.alignment = WRAP
    sh.merge_cells("B6:D6")
    sh.row_dimensions[6].height = 48
    dv = DataValidation(type="list", formula1='"2,1,0"', allow_blank=True)
    sh.add_data_validation(dv)
    dv.add("B4:B5")
    sh["A7"] = "2 — смысл понятен полностью · 1 — частично · 0 — не понятен / пусто / выдумано"
    sh["A7"].font = Font(size=11, italic=True, color="666666")
    # фразы
    sh.append([])
    hdr_row = 9
    for j, t in enumerate(["Время", "large-v3", "Время", "small"], 1):
        c = sh.cell(hdr_row, j, t)
        c.font = FB
        c.fill = HDR
        c.border = BOX
    a, b = segments(n, "large-v3"), segments(n, "small")
    if not a and not b:
        sh.cell(
            hdr_row + 1, 2, "(пусто — модель не распознала речь; проверь, есть ли она в ролике)"
        ).fill = EMPTY
        sh.cell(hdr_row + 1, 4, "(пусто)").fill = EMPTY
    for k in range(max(len(a), len(b))):
        r_ = hdr_row + 1 + k
        ta, xa = a[k] if k < len(a) else ("", "")
        tb, xb = b[k] if k < len(b) else ("", "")
        for j, v in enumerate([ta, xa, tb, xb], 1):
            c = sh.cell(r_, j, v)
            c.font = F
            c.alignment = WRAP
            c.border = BOX
        sh.row_dimensions[r_].height = max(height(xa, 62), height(xb, 62))
    sh.column_dimensions["A"].width = 24
    sh.column_dimensions["B"].width = 62
    sh.column_dimensions["C"].width = 8
    sh.column_dimensions["D"].width = 62
    sh.freeze_panes = "A10"
    # строка в общем листе
    q = f"'{name}'"
    ov.append(
        [
            n,
            L.get("platform"),
            L.get("sound"),
            L.get("topic"),
            "открыть",
            name,
            L.get("duration_s"),
            L.get("words"),
            f'=IF({q}!B4="","",{q}!B4)',
            f'=IF({q}!B5="","",{q}!B5)',
            f'=IF({q}!B6="","",{q}!B6)',
        ]
    )
    ov.cell(i, 5).hyperlink = L.get("url")
    ov.cell(i, 5).font = Font(size=12, color="0563C1", underline="single")
    ov.cell(i, 6).hyperlink = f"#{q}!A1"
    ov.cell(i, 6).font = Font(size=12, color="0563C1", underline="single")
    for j in range(1, 12):
        c = ov.cell(i, j)
        c.border = BOX
        c.alignment = WRAP
        if j not in (5, 6):
            c.font = F
    if not a:
        ov.cell(i, 8).fill = EMPTY
    ov.row_dimensions[i].height = 34
ov.freeze_panes = "A2"
last = len(clips) + 1

# --- Сводка
sm = wb.create_sheet("Сводка", 1)
H, J = f"Оценка!I2:I{last}", f"Оценка!J2:J{last}"
tot_d = sum(by[n]["large-v3"].get("duration_s") or 0 for n in clips)
tot_l = sum(by[n]["large-v3"].get("time_s") or 0 for n in clips)
tot_s = sum(by[n].get("small", {}).get("time_s") or 0 for n in clips)
rows = [
    ["Показатель", "large-v3", "small"],
    ["Оценено роликов", f"=COUNT({H})", f"=COUNT({J})"],
    [
        "Средняя оценка (из 2)",
        f'=IFERROR(ROUND(AVERAGE({H}),2),"")',
        f'=IFERROR(ROUND(AVERAGE({J}),2),"")',
    ],
    ["Смысл понятен полностью (2)", f"=COUNTIF({H},2)", f"=COUNTIF({J},2)"],
    ["Частично (1)", f"=COUNTIF({H},1)", f"=COUNTIF({J},1)"],
    ["Не понятен (0)", f"=COUNTIF({H},0)", f"=COUNTIF({J},0)"],
    [
        "Пустых транскриптов",
        sum(1 for n in clips if not by[n]["large-v3"].get("words")),
        sum(1 for n in clips if not by[n].get("small", {}).get("words")),
    ],
    ["Аудио всего, с", round(tot_d, 1), round(tot_d, 1)],
    ["Время распознавания (GPU T4), с", round(tot_l, 1), round(tot_s, 1)],
    ["RTF (меньше — быстрее)", round(tot_l / tot_d, 3), round(tot_s / tot_d, 3)],
    [],
    ["Вывод (2–3 предложения, для Jira):"],
]
for r in rows:
    sm.append(r)
for c in sm[1]:
    c.font = FB
    c.fill = HDR
for r in sm.iter_rows(min_row=2, max_row=10):
    for c in r:
        c.font = F
sm["A12"].font = FB
sm.merge_cells("A13:C17")
sm["A13"].fill = MAN
sm["A13"].alignment = WRAP
sm["A13"].font = F
sm.column_dimensions["A"].width = 40
sm.column_dimensions["B"].width = 14
sm.column_dimensions["C"].width = 14
sm["A19"] = (
    "RTF = время распознавания / длительность аудио; 0,1 — минута аудио за 6 секунд. Замер: Colab, GPU T4, float16, VAD, beam 5."
)
sm["A19"].font = Font(size=11, italic=True, color="666666")

out = os.path.join(HERE, "asr_check_grading.xlsx")
wb.save(out)
print("ok", len(clips))
